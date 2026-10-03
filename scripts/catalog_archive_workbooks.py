from pathlib import Path
from openpyxl import load_workbook
import json
root=Path('YATA_MPC/references/archive_review')
lines=[]
for p in sorted(root.rglob('*.xlsx')):
 wb=load_workbook(p,read_only=True,data_only=True)
 lines.append('\nWORKBOOK '+str(p.relative_to(root)).encode('unicode_escape').decode())
 for ws in wb.worksheets:
  lines.append(f' SHEET {ws.title.encode("unicode_escape").decode()} rows={ws.max_row} cols={ws.max_column}')
  for row in ws.iter_rows(min_row=1,max_row=6,values_only=True):
   vals=[]
   for i,v in enumerate(row):
    if v is not None: vals.append(f'{i+1}={str(v)[:160].encode("unicode_escape").decode()}')
   if vals: lines.append('  '+' | '.join(vals[:14]))
  if '日射量測定' in p.name:
   for row in ws.iter_rows(min_row=7,max_row=min(15,ws.max_row),values_only=True):
    lines.append('  later '+' | '.join(f'{i+1}={str(v)[:100].encode("unicode_escape").decode()}' for i,v in enumerate(row) if v is not None)[:1500])
Path('YATA_MPC/reports/archive_workbook_catalog_ascii.txt').write_text('\n'.join(lines),encoding='utf-8')
print('workbooks',sum(1 for _ in root.rglob('*.xlsx')))
