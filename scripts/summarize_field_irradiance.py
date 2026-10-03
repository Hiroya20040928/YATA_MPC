from pathlib import Path
from openpyxl import load_workbook
import pandas as pd,numpy as np
p=next(Path('YATA_MPC/references/archive_review').rglob('*日射量測定.xlsx'))
wb=load_workbook(p,read_only=True,data_only=True)
rows=[]
for ws in wb.worksheets:
 meta={}
 for r in range(1,9):
  for row in ws.iter_rows(min_row=r,max_row=r,max_col=2,values_only=True):
   if row[0] is not None: meta[str(row[0])]=row[1]
 vals=[]; times=[]
 for row in ws.iter_rows(min_row=13,max_col=3,values_only=True):
  if row[1] is not None:
   try: vals.append(float(row[1])); times.append(str(row[0]))
   except: pass
 arr=np.asarray(vals)
 rows.append({'sheet':ws.title,'sampling_interval_s':meta.get('記録間隔[s]'),'trigger':str(meta.get('トリガ時刻')),'n':len(vals),'first_time':times[0] if times else None,'last_time':times[-1] if times else None,'mean':float(arr.mean()) if len(arr) else None,'median':float(np.median(arr)) if len(arr) else None,'p10':float(np.quantile(arr,.1)) if len(arr) else None,'p90':float(np.quantile(arr,.9)) if len(arr) else None,'max':float(arr.max()) if len(arr) else None})
df=pd.DataFrame(rows); df.to_csv('YATA_MPC/reports/field_irradiance_logger_summary.csv',index=False)
print(df.to_json(orient='records',force_ascii=True,indent=2))
