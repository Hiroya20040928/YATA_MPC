from pathlib import Path
import fitz,re,json
from docx import Document
root=Path('YATA_MPC').resolve()
arc=root/'references/archive_review'
outdir=root/'references/archive_text'
outdir.mkdir(parents=True,exist_ok=True)
report=[]
for p in sorted(arc.rglob('*')):
    ext=p.suffix.lower()
    if ext=='.pdf':
        doc=fitz.open(p)
        pages=[page.get_text('text') for page in doc]
        text='\n'.join(f'===== PAGE {i+1} =====\n{t}' for i,t in enumerate(pages))
        out=outdir/(p.stem+'.txt')
        out.write_text(text,encoding='utf-8')
        terms=['day5','day 5','maintenance','hatch','anchor','trouble','incident','power','route','gps','control stop','速度','故障','トラブル','ハッチ','充電','日射','距離','アンカー','バッテリー','制御']
        hits={k:len(re.findall(k,text,re.I)) for k in terms if re.search(k,text,re.I)}
        excerpt=[]
        for line in text.splitlines():
            if re.search(r'Day\s*5|Day5|2025/08/28|2025/8/28|\u30cf\u30c3\u30c1|\u30c8\u30e9\u30d6\u30eb|\u30a2\u30f3\u30ab\u30fc|\u30eb\u30fc\u30c8|\u8ddd\u96e2|\u6545\u969c',line,re.I):
                excerpt.append(line.strip())
        report.append({'type':'pdf','file':str(p.relative_to(arc)),'pages':len(doc),'text_chars':len(text),'text_file':str(out.relative_to(root)),'keyword_counts':hits,'day5_excerpt':excerpt[:30]})
    elif ext=='.docx':
        d=Document(p)
        lines=[]
        lines.extend(x.text for x in d.paragraphs if x.text)
        for ti,t in enumerate(d.tables):
            lines.append(f'[TABLE {ti+1}]')
            for row in t.rows:
                lines.append(' | '.join(c.text.replace('\n',' / ') for c in row.cells))
        text='\n'.join(lines)
        out=outdir/(p.stem+'.txt')
        out.write_text(text,encoding='utf-8')
        terms=['Day5','Day 5','maintenance','hatch','anchor','trouble','incident','power','route','GPS','stop','速度','故障','トラブル','ハッチ','距離','アンカー','発電','モーター','消費']
        hits={k:len(re.findall(k,text,re.I)) for k in terms if re.search(k,text,re.I)}
        excerpts=[]
        for line in lines:
            if re.search(r'Day\s*5|Day5|2025/08/28|2025/8/28|\u30cf\u30c3\u30c1|\u30c8\u30e9\u30d6\u30eb|\u30a2\u30f3\u30ab\u30fc|\u30eb\u30fc\u30c8|\u8ddd\u96e2|\u6545\u969c|\u30e2\u30fc\u30bf',line,re.I): excerpts.append(line[:500])
        report.append({'type':'docx','file':str(p.relative_to(arc)),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'text_chars':len(text),'text_file':str(out.relative_to(root)),'keyword_counts':hits,'relevant_excerpt':excerpts[:50]})
(root/'reports'/'archive_document_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print('documents',len(report))
for r in report:
    print(r['type'],r['file'].encode('unicode_escape').decode(), 'pages/paras',r.get('pages',r.get('paragraphs')),'chars',r['text_chars'])
    for s in r.get('day5_excerpt',r.get('relevant_excerpt',[]))[:12]: print(' ',s.encode('unicode_escape').decode())
