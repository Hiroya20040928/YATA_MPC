from pathlib import Path
import zipfile, json
from openpyxl import load_workbook
root=Path('YATA_MPC').resolve()
arc=root/'references/archive_review'
# Catalog all archive spreadsheets and CSV headers/rows.
for p in sorted(arc.rglob('*')):
    if p.suffix.lower() in ('.xlsx','.xlsm'):
        try:
            wb=load_workbook(p,read_only=True,data_only=True)
            print('\nWORKBOOK',p.relative_to(arc))
            for ws in wb.worksheets:
                print(' SHEET',repr(ws.title),'dim',ws.max_row,ws.max_column)
                for row in ws.iter_rows(min_row=1,max_row=min(5,ws.max_row),values_only=True):
                    print('  ',repr(row[:min(12,ws.max_column)]))
        except Exception as e: print('ERROR',p,e)
    elif p.suffix.lower()=='.csv':
        raw=p.read_bytes()
        for enc in ('utf-8-sig','cp932','utf-8'):
            try: text=raw.decode(enc); break
            except UnicodeDecodeError: text=''
        lines=text.splitlines()
        print('\nCSV',p.relative_to(arc),'lines',len(lines),'encoding',enc,'head',repr(lines[:2]))
